# ==============================================================================
# backend/database/exportacion.py
# Lectura de TODOS los datos de un tenant para exportación (respaldo JSON / Excel).
#
# SEGURIDAD:
#   - Todas las funciones reciben `tenant_id`, que SIEMPRE proviene del JWT
#     validado (Depends(get_tenant_id)). Nunca hay un tenant_id controlado
#     por el cliente: aunque alguien manipule la petición, no puede pedir
#     los datos de otro tenant.
#   - Cada consulta filtra por `tenant_id` (aislamiento multitenant).
#
# Formato JSON: envelope versionado para poder restaurar en el futuro.
# Formato XLSX: un libro con una hoja por tabla, para ver los datos en Excel.
# ==============================================================================

import datetime
import json
import uuid
from decimal import Decimal

from database.conexion import query


def _tablas_existentes() -> set[str]:
    """Devuelve el conjunto de tablas del schema 'public' que existen en la DB."""
    filas = query(
        "SELECT table_name FROM information_schema.tables "
        "WHERE table_schema = 'public'"
    )
    return {f["table_name"] for f in filas}


def _parsear_jsonb(valor):
    """psycopg2 devuelve JSONB como string; lo convertimos a objeto real."""
    if valor is None or isinstance(valor, (dict, list)):
        return valor
    try:
        return json.loads(valor)
    except (TypeError, ValueError):
        return valor


def _limpiar_para_json(obj):
    """
    Convierte tipos que no son serializables a JSON (datetimes, UUIDs, Decimals)
    a sus representaciones planas (strings/floats). Aplica a toda la estructura.
    """
    if isinstance(obj, dict):
        return {k: _limpiar_para_json(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_limpiar_para_json(v) for v in obj]
    if isinstance(obj, (datetime.datetime, datetime.date)):
        return obj.isoformat()
    if isinstance(obj, uuid.UUID):
        return str(obj)
    if isinstance(obj, Decimal):
        return float(obj)
    return obj


def get_datos_tenant(tenant_id: str) -> dict:
    """
    Lee TODAS las tablas de negocio del tenant y las devuelve en una estructura
    anidada lista para el JSON (productos traen sus categorías, variaciones,
    recetas e imágenes). Las tablas que no existen en la DB se omiten.
    """
    tablas = _tablas_existentes()
    data: dict = {}

    # ── Productos + relaciones (categorías, variaciones, recetas, imágenes) ──
    if "productos" in tablas:
        productos = query(
            "SELECT id, Producto, Descripcion, Imagen, Estado, codigo_interno, "
            "       codigo_barras, ubicacion, visible_en_catalogo, sufijo_precio, "
            "       fraccionable, tipo_producto, costo_servicio, precio_servicio "
            "FROM productos WHERE tenant_id = %s ORDER BY Producto ASC",
            (tenant_id,)
        )

        categorias: dict = {}
        if "producto_categorias" in tablas and "categorias" in tablas:
            for r in query(
                "SELECT pc.producto_id, c.nombre "
                "FROM producto_categorias pc "
                "JOIN categorias c ON c.id = pc.categoria_id "
                "WHERE c.tenant_id = %s ORDER BY c.nombre ASC",
                (tenant_id,)
            ):
                categorias.setdefault(r["producto_id"], []).append(r["nombre"])

        variaciones: dict = {}
        if "producto_variaciones" in tablas:
            for r in query(
                "SELECT producto_id, id, nombre, precio, foto "
                "FROM producto_variaciones WHERE tenant_id = %s ORDER BY nombre ASC",
                (tenant_id,)
            ):
                variaciones.setdefault(r["producto_id"], []).append({
                    "id": r["id"],
                    "nombre": r["nombre"],
                    "precio": float(r["precio"] or 0),
                    "foto": r.get("foto") or "",
                })

        recetas: dict = {}
        if "producto_recetas" in tablas:
            for r in query(
                "SELECT pr.producto_id, pr.material_id, pr.cantidad, pr.variacion_id, "
                "       m.Producto AS material, v.nombre AS variacion "
                "FROM producto_recetas pr "
                "JOIN productos m ON m.id = pr.material_id "
                "LEFT JOIN producto_variaciones v ON v.id = pr.variacion_id "
                "WHERE pr.tenant_id = %s ORDER BY m.Producto ASC",
                (tenant_id,)
            ):
                recetas.setdefault(r["producto_id"], []).append({
                    "material": r["material"],
                    "cantidad": float(r["cantidad"] or 0),
                    "variacion_id": r.get("variacion_id"),
                    "variacion": r.get("variacion") or "Base",
                })

        imagenes: dict = {}
        if "producto_imagenes" in tablas:
            for r in query(
                "SELECT producto_id, id, url, orden "
                "FROM producto_imagenes WHERE tenant_id = %s ORDER BY orden ASC",
                (tenant_id,)
            ):
                imagenes.setdefault(r["producto_id"], []).append({
                    "id": r["id"],
                    "url": r["url"],
                    "orden": r.get("orden") or 1,
                })

        for p in productos:
            p["categorias"] = categorias.get(p["id"], [])
            p["variaciones"] = variaciones.get(p["id"], [])
            p["recetas"] = recetas.get(p["id"], [])
            p["imagenes"] = imagenes.get(p["id"], [])
        data["productos"] = productos

    # ── Lotes ──
    if "lotes" in tablas:
        data["lotes"] = query(
            "SELECT id, ID_Lote, Producto, producto_id, Costo, Precio_Venta, Stock_Lote, "
            "       Fecha_Entrada, Estado, etiqueta, variacion_id "
            "FROM lotes WHERE tenant_id = %s ORDER BY Producto ASC, Fecha_Entrada ASC",
            (tenant_id,)
        )

    # ── Ventas (consumo JSONB se convierte a objeto) ──
    if "ventas" in tablas:
        ventas = query(
            "SELECT id, n_ticket, Fecha, Producto, producto_id, Cantidad, Precio_Lista, "
            "       Precio_Real, Costo_Unitario, Total_Venta, Ganancia_Bruta, "
            "       Estado, ID_Lote, tipo_producto, variacion, consumo, orden_id "
            "FROM ventas WHERE tenant_id = %s ORDER BY Fecha DESC",
            (tenant_id,)
        )
        for v in ventas:
            v["consumo"] = _parsear_jsonb(v.get("consumo"))
        data["ventas"] = ventas

    # ── Órdenes (tickets; cabecera de cada cobro) ──
    if "ordenes" in tablas:
        data["ordenes"] = query(
            "SELECT id, tenant_id, n_ticket, fecha_ts, total, ganancia, cantidad_items, estado, "
            "       metodo_pago, pagos, propina, monto_recibido, cambio, comision_total "
            "FROM ordenes WHERE tenant_id = %s ORDER BY fecha_ts DESC",
            (tenant_id,)
        )

    # ── Gastos ──
    if "gastos" in tablas:
        data["gastos"] = query(
            "SELECT id, Fecha, Categoria, Descripcion, Monto, Estado, Gasto_Programado_ID "
            "FROM gastos WHERE Tenant_ID = %s ORDER BY Fecha DESC",
            (tenant_id,)
        )

    # ── Categorías de gasto ──
    if "gastos_categorias" in tablas:
        data["gastos_categorias"] = query(
            "SELECT id, nombre FROM gastos_categorias WHERE tenant_id = %s ORDER BY nombre ASC",
            (tenant_id,)
        )

    # ── Gastos programados ──
    if "gastos_programados" in tablas:
        data["gastos_programados"] = query(
            "SELECT id, nombre, tipo, valor, frecuencia, proxima_fecha, "
            "       ultima_ejecucion, created_at "
            "FROM gastos_programados WHERE tenant_id = %s ORDER BY nombre ASC",
            (tenant_id,)
        )

    # ── Categorías de producto ──
    if "categorias" in tablas:
        data["categorias"] = query(
            "SELECT id, nombre, slug, visible_en_catalogo "
            "FROM categorias WHERE tenant_id = %s ORDER BY nombre ASC",
            (tenant_id,)
        )

    # ── Configuración del catálogo público ──
    if "catalogo_config" in tablas:
        filas = query(
            "SELECT * FROM catalogo_config WHERE tenant_id = %s",
            (tenant_id,)
        )
        data["catalogo_config"] = filas[0] if filas else None

    # ── Configuración del negocio (tenant) ──
    if "tenants" in tablas:
        filas = query(
            "SELECT id, logo, modo_precio_sugerido, zona_horaria FROM tenants WHERE id = %s",
            (tenant_id,)
        )
        data["configuracion"] = filas[0] if filas else None

    return _limpiar_para_json(data)


# ==============================================================================
# Vista Excel (XLSX): una hoja por tabla con columnas legibles
# ==============================================================================

def _filas_para_xlsx(data: dict) -> dict[str, list[dict]]:
    """Convierte la estructura de datos en filas planas por hoja de Excel."""
    hojas: dict[str, list[dict]] = {}

    # Productos
    hojas["Productos"] = [
        {
            "Producto": p["producto"],
            "Descripcion": p.get("descripcion") or "",
            "Categorias": ", ".join(p.get("categorias") or []),
            "Tipo": p.get("tipo_producto") or "stock",
            "Precio servicio": p.get("precio_servicio") or 0,
            "Costo servicio": p.get("costo_servicio") or 0,
            "Fraccionable": "Si" if p.get("fraccionable") else "No",
            "Visible en catalogo": "Si" if p.get("visible_en_catalogo") is not False else "No",
            "Codigo interno": p.get("codigo_interno") or "",
            "Codigo de barras": p.get("codigo_barras") or "",
            "Ubicacion": p.get("ubicacion") or "",
            "Estado": p.get("estado") or "",
            "Foto (URL)": p.get("imagen") or "",
        }
        for p in data.get("productos", [])
    ]

    # Variaciones
    hojas["Variaciones"] = []
    for p in data.get("productos", []):
        for v in p.get("variaciones", []):
            hojas["Variaciones"].append({
                "Producto": p["producto"],
                "Variacion": v.get("nombre"),
                "Precio": v.get("precio") or 0,
                "Foto (URL)": v.get("foto") or "",
            })

    # Recetas (compuestos)
    hojas["Recetas"] = []
    for p in data.get("productos", []):
        for r in p.get("recetas", []):
            hojas["Recetas"].append({
                "Compuesto": p["producto"],
                "Material": r.get("material"),
                "Cantidad": r.get("cantidad") or 0,
                "Variacion": r.get("variacion") or "Base",
            })

    # Imágenes extra (galería)
    hojas["Imagenes"] = []
    for p in data.get("productos", []):
        for img in p.get("imagenes", []):
            hojas["Imagenes"].append({
                "Producto": p["producto"],
                "Orden": img.get("orden") or 1,
                "URL": img.get("url"),
            })

    # Lotes
    hojas["Lotes"] = [
        {
            "ID Lote": l.get("id_lote"),
            "Producto": l.get("producto"),
            "Costo": l.get("costo") or 0,
            "Precio venta": l.get("precio_venta") or 0,
            "Stock": l.get("stock_lote") or 0,
            "Fecha entrada": l.get("fecha_entrada") or "",
            "Etiqueta": l.get("etiqueta") or "",
            "Variacion ID": l.get("variacion_id") or "",
            "Estado": l.get("estado") or "",
        }
        for l in data.get("lotes", [])
    ]

    # Ventas
    hojas["Ventas"] = [
        {
            "ID": v.get("id"),
            "Ticket": v.get("n_ticket") or "",
            "Fecha": v.get("fecha") or "",
            "Producto": v.get("producto"),
            "Cantidad": v.get("cantidad") or 0,
            "Precio lista": v.get("precio_lista") or 0,
            "Precio real": v.get("precio_real") or 0,
            "Costo unitario": v.get("costo_unitario") or 0,
            "Total": v.get("total_venta") or 0,
            "Ganancia bruta": v.get("ganancia_bruta") or 0,
            "Tipo": v.get("tipo_producto") or "stock",
            "Variacion": v.get("variacion") or "",
            "Consumo": json.dumps(v.get("consumo"), ensure_ascii=False) if v.get("consumo") else "",
            "ID Lote": v.get("id_lote") or "",
            "Estado": v.get("estado") or "",
        }
        for v in data.get("ventas", [])
    ]

    # Órdenes (tickets)
    hojas["Ordenes"] = [
        {
            "Folio": o.get("n_ticket"),
            "Fecha": str(o.get("fecha_ts") or ""),
            "Total": o.get("total") or 0,
            "Ganancia": o.get("ganancia") or 0,
            "Unidades": o.get("cantidad_items") or 0,
            "Metodo": o.get("metodo_pago") or "",
            "Propina": o.get("propina") or 0,
            "Recibido": o.get("monto_recibido") if o.get("monto_recibido") is not None else "",
            "Cambio": o.get("cambio") if o.get("cambio") is not None else "",
            "Pagos": json.dumps(o.get("pagos"), ensure_ascii=False) if o.get("pagos") else "",
            "Estado": o.get("estado") or "",
        }
        for o in data.get("ordenes", [])
    ]

    # Gastos
    hojas["Gastos"] = [
        {
            "ID": g.get("id"),
            "Fecha": g.get("fecha") or "",
            "Categoria": g.get("categoria") or "",
            "Descripcion": g.get("descripcion") or "",
            "Monto": g.get("monto") or 0,
            "Estado": g.get("estado") or "",
        }
        for g in data.get("gastos", [])
    ]

    # Gastos programados
    hojas["Gastos Programados"] = [
        {
            "Nombre": gp.get("nombre"),
            "Tipo": gp.get("tipo"),
            "Valor": gp.get("valor") or 0,
            "Frecuencia": gp.get("frecuencia"),
            "Proxima fecha": str(gp.get("proxima_fecha") or ""),
            "Ultima ejecucion": str(gp.get("ultima_ejecucion") or ""),
        }
        for gp in data.get("gastos_programados", [])
    ]

    # Categorías de gasto
    hojas["Categorias de Gasto"] = [
        {"Nombre": c.get("nombre")} for c in data.get("gastos_categorias", [])
    ]

    # Categorías de producto
    hojas["Categorias"] = [
        {
            "Nombre": c.get("nombre"),
            "Slug": c.get("slug") or "",
            "Visible en catalogo": "Si" if c.get("visible_en_catalogo") is not False else "No",
        }
        for c in data.get("categorias", [])
    ]

    # Catálogo público (clave/valor)
    cc = data.get("catalogo_config")
    if cc:
        hojas["Catalogo"] = [{"Campo": k, "Valor": v} for k, v in cc.items()]

    # Configuración del negocio (clave/valor)
    conf = data.get("configuracion")
    if conf:
        hojas["Configuracion"] = [{"Campo": k, "Valor": v} for k, v in conf.items()]

    return hojas


def generar_xlsx(tenant_id: str) -> bytes:
    """
    Genera el archivo XLSX (bytes) con una hoja por tabla de negocio del tenant.
    Los datos vienen de get_datos_tenant() — siempre filtrados por tenant_id.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    import io

    datos = get_datos_tenant(tenant_id)
    hojas = _filas_para_xlsx(datos)

    wb = Workbook()
    wb.remove(wb.active)  # quitar la hoja por defecto

    for nombre, filas in hojas.items():
        ws = wb.create_sheet(title=nombre[:31])
        if not filas:
            ws.append(["Sin datos"])
            continue
        headers = list(filas[0].keys())
        ws.append(headers)
        for fila in filas:
            ws.append([fila.get(h) for h in headers])
        # Encabezado en negrita
        for col in range(1, len(headers) + 1):
            ws.cell(row=1, column=col).font = Font(bold=True)
        # Ancho de columna aproximado (máx 60 para no desbordar)
        for col in range(1, len(headers) + 1):
            ancho = 12
            try:
                max_len = max(len(str(fila.get(headers[col - 1], ""))) for fila in filas)
                ancho = max(12, min(max_len + 2, 60))
            except Exception:
                pass
            ws.column_dimensions[get_column_letter(col)].width = ancho

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
